"""
Script to convert Markdown manual book to DOCX format
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re

def create_manual_book_docx(md_file, output_file):
    """Convert markdown file to formatted DOCX"""
    
    # Create document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Read markdown file
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Process line by line
    lines = content.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Skip empty lines (but add spacing)
        if not line:
            doc.add_paragraph()
            i += 1
            continue
        
        # Handle headers
        if line.startswith('# '):
            # H1 - Main title or chapter
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        elif line.startswith('## '):
            # H2 - Section
            doc.add_heading(line[3:], level=2)
            
        elif line.startswith('### '):
            # H3 - Subsection
            doc.add_heading(line[4:], level=3)
            
        # Handle tables
        elif line.startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|'):
            # Parse table
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            i -= 1
            
            # Skip separator line
            if len(table_lines) > 1 and '---' in table_lines[1]:
                table_lines.pop(1)
            
            # Create table
            if table_lines:
                # Parse first row to get column count
                cols = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                rows = len(table_lines)
                
                table = doc.add_table(rows=rows, cols=len(cols))
                table.style = 'Light Grid Accent 1'
                
                for row_idx, table_line in enumerate(table_lines):
                    cells = [c.strip() for c in table_line.split('|')[1:-1]]
                    for col_idx, cell_content in enumerate(cells):
                        if col_idx < len(cols):
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = cell_content
                            # Bold header row
                            if row_idx == 0:
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.font.bold = True
                
                doc.add_paragraph()
        
        # Handle code blocks
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            
            # Add code block
            if code_lines:
                p = doc.add_paragraph('\n'.join(code_lines))
                p.style = 'No Spacing'
                for run in p.runs:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                doc.add_paragraph()
        
        # Handle placeholders
        elif '[PLACEHOLDER' in line:
            p = doc.add_paragraph(line)
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 0, 0)
        
        # Handle bold text
        elif line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph()
            run = p.add_run(line[2:-2])
            run.font.bold = True
        
        # Handle horizontal rules
        elif line == '---':
            doc.add_paragraph('_' * 50)
        
        # Regular paragraphs
        else:
            # Check for inline formatting
            p = doc.add_paragraph()
            
            # Split by bold markers
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                elif part:
                    p.add_run(part)
        
        i += 1
    
    # Save document
    doc.save(output_file)
    print(f"Document saved successfully: {output_file}")

if __name__ == "__main__":
    md_file = "MANUAL_BOOK_DRAFT.md"
    output_file = "MANUAL_BOOK_APLIKASI_GAJI_KARYAWAN_PT_SINTORY.docx"
    
    try:
        create_manual_book_docx(md_file, output_file)
    except Exception as e:
        print(f"Error: {e}")
        print("\nTrying alternative method using pandoc...")
        import subprocess
        try:
            subprocess.run([
                "pandoc",
                md_file,
                "-o", output_file,
                "--toc"
            ], check=True)
            print(f"Document converted successfully using pandoc: {output_file}")
        except:
            print("Pandoc not available. Please install python-docx or pandoc.")
